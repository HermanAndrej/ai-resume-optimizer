"""Render an ExportResume to ATS-safe DOCX bytes.

Hard-coded styling: Calibri 11pt body, 14pt bold section headers, 18pt bold
name, 0.75" margins, single column, no tables for layout, no images, no
headers/footers. Sections are omitted when empty.
"""
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from backend.models import (
    CertificationEntry,
    EducationEntry,
    ExportExperience,
    ExportResume,
    ExportSkillGroup,
    PersonalInfo,
    ProjectEntry,
)

BODY_FONT = "Calibri"
BODY_SIZE_PT = 11
NAME_SIZE_PT = 18
HEADING_SIZE_PT = 14
MARGIN_INCHES = 0.75


def _date_range(start: str, end: str) -> str:
    if start and end:
        return f"{start} – {end}"
    return start or end or ""


def _set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(BODY_SIZE_PT)


def _set_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(MARGIN_INCHES)
        section.bottom_margin = Inches(MARGIN_INCHES)
        section.left_margin = Inches(MARGIN_INCHES)
        section.right_margin = Inches(MARGIN_INCHES)


def _render_personal(doc: Document, personal: PersonalInfo) -> None:
    if personal.full_name:
        p = doc.add_paragraph()
        run = p.add_run(personal.full_name)
        run.bold = True
        run.font.size = Pt(NAME_SIZE_PT)
        p.paragraph_format.space_after = Pt(2)

    contact_parts = [c for c in (personal.email, personal.phone, personal.location) if c]
    if contact_parts:
        p = doc.add_paragraph()
        run = p.add_run(" · ".join(contact_parts))
        run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(8)


def _render_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(HEADING_SIZE_PT)


def _render_summary(doc: Document, summary: str) -> None:
    if not summary:
        return
    _render_section_heading(doc, "Summary")
    p = doc.add_paragraph(summary)
    p.paragraph_format.space_after = Pt(4)


def _render_experience(doc: Document, items: list[ExportExperience]) -> None:
    if not items:
        return
    _render_section_heading(doc, "Experience")
    for exp in items:
        # Header line: Title — Company  ·  dates · location
        p = doc.add_paragraph()
        title_run = p.add_run(exp.title or "")
        title_run.bold = True
        if exp.company:
            company_run = p.add_run(f" — {exp.company}")
            company_run.bold = True
        meta_parts = [_date_range(exp.start_date, exp.end_date), exp.location]
        meta = " · ".join([m for m in meta_parts if m])
        if meta:
            meta_run = p.add_run(f"  ·  {meta}")
            meta_run.italic = True
            meta_run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)

        for bullet_text in exp.bullets:
            if bullet_text:
                bp = doc.add_paragraph(bullet_text, style="List Bullet")
                bp.paragraph_format.space_after = Pt(0)

        # Trailing space after each role
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)


def _render_skills(doc: Document, groups: list[ExportSkillGroup]) -> None:
    if not groups:
        return
    _render_section_heading(doc, "Skills")
    for group in groups:
        if not group.items:
            continue
        p = doc.add_paragraph()
        cat_run = p.add_run(f"{group.category}: ")
        cat_run.bold = True
        p.add_run(", ".join(group.items))
        p.paragraph_format.space_after = Pt(2)


def _render_projects(doc: Document, items: list[ProjectEntry]) -> None:
    if not items:
        return
    _render_section_heading(doc, "Projects")
    for proj in items:
        if not (proj.name or proj.description):
            continue
        p = doc.add_paragraph()
        if proj.name:
            name_run = p.add_run(proj.name)
            name_run.bold = True
        if proj.tech_stack:
            tech_run = p.add_run(f"  ·  {proj.tech_stack}")
            tech_run.italic = True
            tech_run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)
        if proj.description:
            desc = doc.add_paragraph(proj.description)
            desc.paragraph_format.space_after = Pt(2)
        if proj.bullets:
            for line in proj.bullets.splitlines():
                line = line.strip().lstrip("-•· ").strip()
                if line:
                    bp = doc.add_paragraph(line, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(0)


def _render_education(doc: Document, items: list[EducationEntry]) -> None:
    if not items:
        return
    _render_section_heading(doc, "Education")
    for edu in items:
        if not (edu.institution or edu.degree):
            continue
        p = doc.add_paragraph()
        if edu.institution:
            inst_run = p.add_run(edu.institution)
            inst_run.bold = True
        meta_parts = []
        if edu.degree:
            meta_parts.append(edu.degree + (f", {edu.field}" if edu.field else ""))
        elif edu.field:
            meta_parts.append(edu.field)
        date_str = _date_range(edu.start_date, edu.end_date)
        if date_str:
            meta_parts.append(date_str)
        if edu.gpa:
            meta_parts.append(f"GPA {edu.gpa}")
        if meta_parts:
            meta_run = p.add_run(f"  ·  {' · '.join(meta_parts)}")
            meta_run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        if edu.highlights:
            hl = doc.add_paragraph(edu.highlights)
            hl.paragraph_format.space_after = Pt(2)


def _render_certifications(doc: Document, items: list[CertificationEntry]) -> None:
    if not items:
        return
    _render_section_heading(doc, "Certifications")
    for cert in items:
        if not cert.name:
            continue
        p = doc.add_paragraph()
        name_run = p.add_run(cert.name)
        name_run.bold = True
        meta_parts = [cert.issuer, cert.date]
        meta = " · ".join([m for m in meta_parts if m])
        if meta:
            p.add_run(f"  ·  {meta}").font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)


def render_docx(resume: ExportResume) -> bytes:
    """Render the export resume to .docx bytes."""
    doc = Document()
    _set_default_font(doc)
    _set_margins(doc)

    _render_personal(doc, resume.personal)
    _render_summary(doc, resume.summary)
    _render_experience(doc, resume.experience)
    _render_skills(doc, resume.skills_grouped)
    _render_projects(doc, resume.projects)
    _render_education(doc, resume.education)
    _render_certifications(doc, resume.certifications)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
