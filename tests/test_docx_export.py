"""Unit tests for docx_export.render_docx — DOCX rendering."""
from io import BytesIO

from docx import Document

from backend.models import (
    CertificationEntry,
    EducationEntry,
    ExportExperience,
    ExportResume,
    ExportSkillGroup,
    PersonalInfo,
    ProjectEntry,
)
from backend.services.docx_export import render_docx


def _populated_resume() -> ExportResume:
    return ExportResume(
        personal=PersonalInfo(
            full_name="Jane Doe",
            email="jane@example.com",
            phone="555-1212",
            location="Berlin",
        ),
        summary="Senior backend engineer with 8 years of Python.",
        experience=[
            ExportExperience(
                company="Acme Corp",
                title="Senior Engineer",
                location="Remote",
                start_date="2020",
                end_date="Present",
                bullets=[
                    "Reduced p99 latency by 40%",
                    "Migrated 12 services to Kubernetes",
                ],
            ),
            ExportExperience(
                company="Old Job",
                title="Junior Dev",
                start_date="2018",
                end_date="2020",
                bullets=["Wrote tests"],
            ),
        ],
        skills_grouped=[
            ExportSkillGroup(category="Languages", items=["Python", "Go"]),
            ExportSkillGroup(category="Cloud", items=["Kubernetes", "AWS"]),
        ],
        projects=[
            ProjectEntry(
                name="logsift",
                description="CLI for log analysis",
                tech_stack="Python",
            ),
        ],
        education=[
            EducationEntry(
                institution="TU Berlin",
                degree="BSc",
                field="Computer Science",
                end_date="2018",
            ),
        ],
        certifications=[
            CertificationEntry(name="AWS SAA", issuer="Amazon", date="2022"),
        ],
    )


def _open(bytes_: bytes) -> Document:
    return Document(BytesIO(bytes_))


def _all_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


class TestRenderDocxValid:
    def test_returns_bytes_with_zip_magic(self):
        out = render_docx(_populated_resume())
        assert isinstance(out, bytes)
        # All DOCX files are zip archives → start with PK\x03\x04
        assert out.startswith(b"PK\x03\x04")

    def test_empty_resume_renders_valid_docx(self):
        out = render_docx(ExportResume())
        assert out.startswith(b"PK\x03\x04")
        # And reads as a valid Document
        doc = _open(out)
        assert doc is not None


class TestPopulatedContent:
    def test_personal_info_present(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Jane Doe" in text
        assert "jane@example.com" in text
        assert "555-1212" in text
        assert "Berlin" in text

    def test_summary_section_present(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Summary" in text
        assert "8 years of Python" in text

    def test_experience_section_with_all_entries(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Experience" in text
        assert "Senior Engineer" in text
        assert "Acme Corp" in text
        assert "Junior Dev" in text
        assert "Old Job" in text

    def test_experience_bullets_rendered(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Reduced p99 latency by 40%" in text
        assert "Migrated 12 services to Kubernetes" in text

    def test_skills_grouped_by_category(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Languages:" in text
        assert "Python, Go" in text
        assert "Cloud:" in text
        assert "Kubernetes, AWS" in text

    def test_projects_education_certifications_present(self):
        doc = _open(render_docx(_populated_resume()))
        text = _all_text(doc)
        assert "Projects" in text
        assert "logsift" in text
        assert "Education" in text
        assert "TU Berlin" in text
        assert "Certifications" in text
        assert "AWS SAA" in text


class TestEmptySections:
    def test_no_summary_section_when_summary_empty(self):
        resume = _populated_resume()
        resume.summary = ""
        doc = _open(render_docx(resume))
        text = _all_text(doc)
        assert "Summary" not in text

    def test_no_education_section_when_empty(self):
        resume = _populated_resume()
        resume.education = []
        doc = _open(render_docx(resume))
        text = _all_text(doc)
        assert "Education" not in text
        # but the rest should still be present
        assert "Experience" in text

    def test_no_certifications_when_empty(self):
        resume = _populated_resume()
        resume.certifications = []
        doc = _open(render_docx(resume))
        text = _all_text(doc)
        assert "Certifications" not in text

    def test_no_projects_when_empty(self):
        resume = _populated_resume()
        resume.projects = []
        doc = _open(render_docx(resume))
        text = _all_text(doc)
        assert "Projects" not in text


class TestStyling:
    def test_default_font_is_calibri(self):
        out = render_docx(_populated_resume())
        doc = _open(out)
        normal = doc.styles["Normal"]
        assert normal.font.name == "Calibri"

    def test_default_body_size_is_11pt(self):
        from docx.shared import Pt
        out = render_docx(_populated_resume())
        doc = _open(out)
        normal = doc.styles["Normal"]
        assert normal.font.size == Pt(11)

    def test_margins_are_three_quarters_inch(self):
        from docx.shared import Inches
        out = render_docx(_populated_resume())
        doc = _open(out)
        section = doc.sections[0]
        assert section.top_margin == Inches(0.75)
        assert section.bottom_margin == Inches(0.75)
        assert section.left_margin == Inches(0.75)
        assert section.right_margin == Inches(0.75)

    def test_no_tables_for_layout(self):
        """ATS-safe — no tables in the document."""
        out = render_docx(_populated_resume())
        doc = _open(out)
        assert len(doc.tables) == 0

    def test_no_inline_images(self):
        out = render_docx(_populated_resume())
        doc = _open(out)
        # python-docx exposes inline shapes on the document
        assert len(doc.inline_shapes) == 0


class TestDateFormatting:
    def test_experience_date_range_renders(self):
        resume = ExportResume(
            experience=[
                ExportExperience(
                    company="X", title="Y", start_date="2020", end_date="2023",
                ),
            ],
        )
        text = _all_text(_open(render_docx(resume)))
        assert "2020" in text
        assert "2023" in text

    def test_only_end_date_renders(self):
        resume = ExportResume(
            experience=[
                ExportExperience(company="X", title="Y", end_date="Present"),
            ],
        )
        text = _all_text(_open(render_docx(resume)))
        assert "Present" in text
