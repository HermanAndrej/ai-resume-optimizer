"""Tests for resume text extraction (DOCX round-trip; PDF skipped by default)."""
import io

import pytest

from backend.services.extraction import extract_docx, extract_text


def _make_docx(paragraphs: list[str], table_cells: list[str] | None = None) -> bytes:
    """Build a minimal in-memory DOCX for testing."""
    import docx

    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)

    if table_cells:
        table = doc.add_table(rows=1, cols=len(table_cells))
        for i, cell_text in enumerate(table_cells):
            table.cell(0, i).text = cell_text

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestExtractDocx:
    def test_basic_paragraphs(self):
        data = _make_docx(["Hello World", "Second paragraph"])
        result = extract_docx(data)
        assert "Hello World" in result
        assert "Second paragraph" in result

    def test_table_cells_included(self):
        data = _make_docx(["Intro"], table_cells=["Python", "FastAPI", "SQLite"])
        result = extract_docx(data)
        assert "Python" in result
        assert "FastAPI" in result
        assert "SQLite" in result

    def test_empty_paragraphs_skipped(self):
        data = _make_docx(["Real content", "", "   "])
        result = extract_docx(data)
        lines = [l for l in result.split("\n") if l]
        assert lines == ["Real content"]

    def test_returns_string(self):
        data = _make_docx(["Test"])
        assert isinstance(extract_docx(data), str)


class TestExtractText:
    def test_docx_dispatches_correctly(self):
        data = _make_docx(["Alice Smith", "Software Engineer", "alice@example.com"] * 5)
        result = extract_text("resume.docx", data)
        assert "Alice Smith" in result

    def test_docx_uppercase_extension(self):
        data = _make_docx(["Alice Smith", "Software Engineer at Google"] * 5)
        result = extract_text("resume.DOCX", data)
        assert "Alice" in result

    def test_unsupported_extension_raises(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text("resume.txt", b"some text content here")

    def test_rtf_raises(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text("resume.rtf", b"rtf content")

    def test_too_short_raises(self):
        # A DOCX with very little text should raise the "too short" error.
        data = _make_docx(["Hi"])
        with pytest.raises(ValueError, match="too short"):
            extract_text("resume.docx", data)

    def test_sufficient_text_passes(self):
        # 50+ chars of content should succeed.
        long_content = ["Alice Smith — Senior Software Engineer at Acme Corp"] * 3
        data = _make_docx(long_content)
        result = extract_text("resume.docx", data)
        assert len(result.strip()) >= 50


@pytest.mark.skip(reason="PDF fixture not generated in CI — test manually with a real PDF")
class TestExtractPdf:
    def test_pdf_round_trip(self):
        # To run: provide a real text-based PDF and call extract_text("resume.pdf", data)
        pass
